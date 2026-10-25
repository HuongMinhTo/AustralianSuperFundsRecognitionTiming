"""Word and Excel report for the asymmetric-recognition test (run after asymmetry_model.py).
Output: output/asymmetry/Asymmetric_recognition_results.docx, asymmetric_recognition_tables.xlsx, fig1_recognition_curves_pos_neg.png
Run from anywhere:  python3 code/6_asymmetry/report_asymmetry.py"""
import os
import copy
import subprocess
import tempfile
import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(ROOT, 'output', 'asymmetry')
os.makedirs(OUT, exist_ok=True)
R = pd.read_pickle(os.path.join(ROOT, 'output', 'estimates', 'asymmetry', 'asym_results.pkl'))
GROUPS = [(0, 0), (1, 1), (2, 4), (5, 8), (9, 13), (14, 26)]
NG = np.array([b - a + 1 for a, b in GROUPS])
WN = {'equal': 'Equal-weighted', 'aum': 'AUM-weighted'}
CASES = [(('own', 'A', False), 'Sign of own news, Model A (main)'),
         (('own', 'B', False), 'Sign of own news, Model B (week fixed effects)'),
         (('own', 'A', True), 'Sign of own news, Model A, excl. Feb-Jun 2020'),
         (('mkt', 'A', False), 'Sign of the equity market, Model A'),
         (('large', 'A', False), 'Largest 10% falls vs largest 10% rises, Model A')]
MAIN = ('own', 'A', False)


def p2(v):
    return 2 * (1 - stats.norm.cdf(abs(v[0] / v[1])))


def cell(v, d=3):
    return f'{v[0]:.{d}f} ({v[1]:.{d}f})'


def star(p):
    return '***' if p < 0.01 else '**' if p < 0.05 else '*' if p < 0.10 else ''


def curve(b):
    beta = np.concatenate([np.full(n, x) for n, x in zip(NG, b)])
    return np.cumsum(beta) / beta.sum()


# ------------------------------------------------------------------ tables
LAB = {('XU', 'p'): 'Unlisted, positive', ('XU', 'n'): 'Unlisted, negative',
       ('XL', 'p'): 'Listed, positive', ('XL', 'n'): 'Listed, negative'}
ROWS = [('b0', 'Same-week response'), ('D', 'Delayed response, weeks 1-26 (D)'), ('T', 'Total response, weeks 0-26'),
        ('share', 'Delayed share (D / Total)'), ('R1', 'Share recognised by week 1'), ('R4', 'Share recognised by week 4'),
        ('R13', 'Share recognised by week 13')]

# Table 1: how much news falls in each state
t1 = []
for (var, lab) in [('own', 'Sign of own news'), ('mkt', 'Sign of the equity market'), ('large', 'Largest 10% falls / rises')]:
    c = R[('equal', var, 'A', False)]['counts']
    t1.append({'Definition of the state': lab,
               'Unlisted news, positive: option-weeks (weeks)': f"{c[('XU', 'p')]['option_weeks']:,} ({c[('XU', 'p')]['weeks']})",
               'Unlisted news, negative: option-weeks (weeks)': f"{c[('XU', 'n')]['option_weeks']:,} ({c[('XU', 'n')]['weeks']})",
               'Listed news, positive: option-weeks (weeks)': f"{c[('XL', 'p')]['option_weeks']:,} ({c[('XL', 'p')]['weeks']})",
               'Listed news, negative: option-weeks (weeks)': f"{c[('XL', 'n')]['option_weeks']:,} ({c[('XL', 'n')]['weeks']})"})
t1 = pd.DataFrame(t1)
t1_note = ('Equal-weighted estimation sample (82,305 option-weeks, 220 options, July 2015 to August 2026). Option-weeks with '
           'non-zero news in each state; distinct weeks in parentheses. Own news: the sign of the option\'s unlisted news X^U '
           '(for unlisted news) or listed news X^L (for listed news) in that week. Equity market: 0.5 ASX 300 + 0.5 world '
           'ex-Australia ETF return over the option\'s return window. Largest falls / rises: weeks in the bottom / top 10% of '
           f'that market return (below {100 * R["P10"]:.1f}% / above {100 * R["P90"]:.1f}% a week); the middle 80% is a '
           'separate third state. Options with no unlisted allocation contribute zero unlisted news.')

# Table 2: main estimates
t2 = []
for wgt in ('equal', 'aum'):
    a = R[(wgt,) + MAIN]
    t2.append({'Measure': f'Panel {"A" if wgt == "equal" else "B"}: {WN[wgt]} ({a["n"]:,} option-weeks)',
               **{LAB[k]: '' for k in LAB}})
    for k_, lab in ROWS:
        t2.append({'Measure': lab, **{LAB[k]: cell(a['parts'][k][k_]) for k in LAB}})
t2 = pd.DataFrame(t2)
t2_note = ('Model A (investment-option fixed effects), news split by the sign of the option\'s own news in each week. '
           'Responses are sums of the weekly lag coefficients (grouped lags 0, 1, 2-4, 5-8, 9-13, 14-26). Shares are ratios '
           'to the total response through week 26. Driscoll-Kraay standard errors (Bartlett kernel, bandwidth 8) in '
           'parentheses; standard errors of ratios by the delta method. Defensive and other exposures are controls.')

# Table 3: tests across specifications
t3, raw3 = [], []
for wgt in ('equal', 'aum'):
    for case, lab in CASES:
        t = R[(wgt,) + case]['tests']
        pw = 1 - stats.chi2.cdf(t['XU_wald'], 6)
        pwl = 1 - stats.chi2.cdf(t['XL_wald'], 6)
        t3.append({'Specification': f'{WN[wgt]}: {lab}',
                   'Unlisted: delayed share, negative minus positive': cell(t['XU_dshare']) + star(p2(t['XU_dshare'])),
                   'Listed: delayed share, negative minus positive': cell(t['XL_dshare']) + star(p2(t['XL_dshare'])),
                   'Difference-in-differences (unlisted minus listed)': cell(t['did_share']) + star(p2(t['did_share'])),
                   'Unlisted: share by week 1, negative minus positive': cell(t['XU_dR1']) + star(p2(t['XU_dR1'])),
                   'Unlisted: same lag profile, Wald p-value': f'{pw:.3f}'})
        raw3.append(dict(weighting=wgt, specification=lab, **{f'{k}_est': v[0] for k, v in t.items() if not k.endswith('wald')},
                         **{f'{k}_se': v[1] for k, v in t.items() if not k.endswith('wald')},
                         **{f'{k}_p': p2(v) for k, v in t.items() if not k.endswith('wald')},
                         XU_wald=t['XU_wald'], XU_wald_p=pw, XL_wald=t['XL_wald'], XL_wald_p=pwl,
                         n=R[(wgt,) + case]['n']))
t3 = pd.DataFrame(t3)
raw3 = pd.DataFrame(raw3)
t3_note = ('Delayed share = delayed response (weeks 1-26) / total response. A positive difference means that negative news '
           'is recognised more slowly than positive news. The difference-in-differences nets out any asymmetry that also '
           'appears in listed news, which is priced at market every day. Wald test: the six lag-group coefficients are the '
           'same for positive and negative unlisted news (6 degrees of freedom). Model B adds week fixed effects. Driscoll-'
           'Kraay standard errors; two-sided tests: * p < 0.10, ** p < 0.05, *** p < 0.01.')

# ------------------------------------------------------------------ figure
fig, axs = plt.subplots(1, 2, figsize=(10, 4), sharey=True)
h = np.arange(27)
for ax, wgt in zip(axs, ('equal', 'aum')):
    P = R[(wgt,) + MAIN]['parts']
    ax.plot(h, curve(P[('XU', 'p')]['b']), color='black', lw=1.8, label='Unlisted, positive news')
    ax.plot(h, curve(P[('XU', 'n')]['b']), color='black', lw=1.8, ls='--', label='Unlisted, negative news')
    ax.plot(h, curve(P[('XL', 'p')]['b']), color='0.55', lw=1.4, label='Listed, positive news')
    ax.plot(h, curve(P[('XL', 'n')]['b']), color='0.55', lw=1.4, ls='--', label='Listed, negative news')
    ax.axhline(1, color='0.8', lw=0.8)
    ax.set_title(WN[wgt], fontsize=11, family='serif')
    ax.set_xlabel('Weeks after the news (h)', family='serif')
    ax.set_xlim(0, 26); ax.set_xticks([0, 1, 4, 8, 13, 26])
    for s in ('top', 'right'):
        ax.spines[s].set_visible(False)
axs[0].set_ylabel('Share of total response recognised', family='serif')
axs[0].legend(frameon=False, fontsize=8.5, loc='lower right', prop={'family': 'serif', 'size': 8.5})
fig.tight_layout()
FIG = os.path.join(OUT, 'fig1_recognition_curves_pos_neg.png')
fig.savefig(FIG, dpi=200)
plt.close(fig)

# ------------------------------------------------------------------ key numbers for the text
eA, wA = R[('equal',) + MAIN], R[('aum',) + MAIN]
eT, wT = eA['tests'], wA['tests']
eP, wP = eA['parts'], wA['parts']
eB, wB = R[('equal', 'own', 'B', False)]['tests'], R[('aum', 'own', 'B', False)]['tests']
pvals_share = [p2(R[(w_,) + c]['tests']['XU_dshare']) for w_ in ('equal', 'aum') for c, _ in CASES]
pvals_did = [p2(R[(w_,) + c]['tests']['did_share']) for w_ in ('equal', 'aum') for c, _ in CASES]
se_share = [R[(w_,) + c]['tests']['XU_dshare'][1] for w_ in ('equal', 'aum') for c, _ in CASES]
MDE = 2.8 * np.median(se_share)

# ------------------------------------------------------------------ Word
doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for a_ in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(s, a_, Cm(2.3))
for st in doc.styles:
    try:
        st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a_ in list(rf.attrib):
            del rf.attrib[a_]
        for a_ in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a_), 'Times New Roman')
    except Exception:
        pass


def style(name):
    for s_ in doc.styles:
        if s_.name == name:
            return s_
    return None


for el in doc.styles.element.iter(qn('w:color')):
    for a_ in ('w:themeColor', 'w:themeTint', 'w:themeShade'):
        if el.get(qn(a_)) is not None:
            del el.attrib[qn(a_)]
    el.set(qn('w:val'), '000000')
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
style('Normal').font.size = Pt(11)
for h_, sz in (('Title', 16), ('Heading 1', 13), ('Heading 2', 11.5)):
    s_ = style(h_)
    if s_ is not None:
        s_.font.size = Pt(sz); s_.font.bold = True


def border(c, edge, sz):
    tcPr = c._tc.get_or_add_tcPr(); b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), '000000')
    b.append(e)


def add_table(num, title, t, note, widths, fs=8.5, panel_rows=()):
    p_ = doc.add_paragraph(); r_ = p_.add_run(f'{num}. {title}'); r_.bold = True
    p_.paragraph_format.space_before = Pt(12); p_.paragraph_format.keep_with_next = True
    tb = doc.add_table(rows=len(t) + 1, cols=t.shape[1]); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for gc, w_ in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w_ * 567)))
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tb._tbl.tblPr.append(lay)
    for j, w_ in enumerate(widths):
        for c in tb.columns[j].cells:
            c.width = Cm(w_)
    for j, c in enumerate(t.columns):
        tb.cell(0, j).text = str(c)
    for i, row in enumerate(t.itertuples(index=False)):
        for j, v in enumerate(row):
            tb.cell(i + 1, j).text = str(v)
    n = len(t) + 1
    for i in range(n):
        for j in range(t.shape[1]):
            c = tb.cell(i, j)
            if i == 0:
                border(c, 'top', 12); border(c, 'bottom', 6)
            if i == n - 1:
                border(c, 'bottom', 12)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.space_before = Pt(1)
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                pp.paragraph_format.keep_with_next = i < n - 1
                for rr in pp.runs:
                    rr.font.size = Pt(fs); rr.font.bold = (i == 0) or (i - 1 in panel_rows)
                    rr.font.italic = (i - 1 in panel_rows)
    q = doc.add_paragraph(); rn = q.add_run('Note: ' + note); rn.font.size = Pt(9)


def add_figure(num, title, path, note):
    p_ = doc.add_paragraph(); r_ = p_.add_run(f'{num}. {title}'); r_.bold = True
    p_.paragraph_format.space_before = Pt(12); p_.paragraph_format.keep_with_next = True
    doc.add_picture(path, width=Cm(16)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    q = doc.add_paragraph(); rn = q.add_run('Note: ' + note); rn.font.size = Pt(9)


def md_block(md):
    """markdown with LaTeX maths -> Word paragraphs (maths as native equation-editor objects) via pandoc"""
    with tempfile.TemporaryDirectory() as td:
        src, dst = os.path.join(td, 'x.md'), os.path.join(td, 'x.docx')
        open(src, 'w').write(md)
        subprocess.run(['pandoc', src, '-f', 'markdown', '-t', 'docx', '-o', dst], check=True)
        body = Document(dst).element.body
        for el in list(body):
            if el.tag == qn('w:sectPr'):
                continue
            el = copy.deepcopy(el)
            for ps in el.iter(qn('w:pStyle')):
                ps.getparent().remove(ps)
            for pp in ([el] if el.tag == qn('w:p') else []):
                ppr = pp.find(qn('w:pPr'))
                if ppr is None:
                    ppr = OxmlElement('w:pPr'); pp.insert(0, ppr)
                sp = OxmlElement('w:spacing'); sp.set(qn('w:after'), '120'); ppr.append(sp)
            doc.element.body.insert(len(doc.element.body) - 1, el)


doc.add_heading('Is negative market news recognised more slowly than positive news? Asymmetric recognition in '
                'unlisted exposures', level=0)
doc.add_paragraph(f'Minh Huong · 23 funds, 220 investment options, weekly returns July 2015 to August 2026 · '
                  'equal-weighted and AUM-weighted estimates')

doc.add_heading('1. Question', level=1)
md_block(r"""
The main model shows that market news relevant to unlisted assets is recognised in unit prices over several weeks. This report asks whether that delay depends on the direction of the news. Two arguments point in opposite directions. If valuers or managers are reluctant to mark assets down, negative news would be recognised more slowly than positive news, as Bollen and Pool (2008) find for losses reported by hedge funds. If instead trustees revalue unlisted assets out of cycle after large market falls, as APRA expects (APRA, 2024b), negative news would be recognised faster. The hypothesis is therefore two-sided:

**H7: Asymmetric recognition.** The speed with which unlisted information is recognised differs between positive and negative market news.
""")

doc.add_heading('2. Method', level=1)
md_block(r"""
Each week's exposure-weighted news is split by the state of the market in that week. In the main specification the state is the sign of the option's own news:

$$X^{U+}_{i,t}=X^{U}_{i,t}\,\mathbf{1}\!\left[X^{U}_{i,t}>0\right],\qquad X^{U-}_{i,t}=X^{U}_{i,t}\,\mathbf{1}\!\left[X^{U}_{i,t}\le 0\right],$$

and the listed news $X^{L}_{i,t}$ is split in the same way by its own sign. The grouped lags of each part, $Z^{U\pm}_{i,g,t}=\sum_{k\in g}X^{U\pm}_{i,t-k}$ and $Z^{L\pm}_{i,g,t}=\sum_{k\in g}X^{L\pm}_{i,t-k}$, replace the pooled news in the main model:

$$r_{i,t}=\alpha_{i}+\delta_{t}+\sum_{g\in G}\left(\beta^{U+}_{g}Z^{U+}_{i,g,t}+\beta^{U-}_{g}Z^{U-}_{i,g,t}+\beta^{L+}_{g}Z^{L+}_{i,g,t}+\beta^{L-}_{g}Z^{L-}_{i,g,t}\right)+\Gamma'W_{i,t}+\varepsilon_{i,t},$$

where $G=\{0,1,2\!:\!4,5\!:\!8,9\!:\!13,14\!:\!26\}$, $\delta_{t}=0$ in Model A, and the other terms are as in the main model. If recognition is symmetric, $\beta^{U+}_{g}=\beta^{U-}_{g}$ for every group.

Negative weeks are fewer and larger than positive weeks, so the size of the responses is not comparable between the two. Speed is therefore compared with a scale-free measure, the delayed share of each response,

$$s^{U\pm}=\frac{D^{U\pm}}{Total^{U\pm}}=\frac{\sum_{g\neq\{0\}}n_{g}\beta^{U\pm}_{g}}{\sum_{g}n_{g}\beta^{U\pm}_{g}},$$

and the asymmetry is $A^{U}=s^{U-}-s^{U+}$; a positive value means that negative news is recognised more slowly. Listed assets are priced at market every day, so their response should be symmetric; any asymmetry $A^{L}$ in listed news reflects the method or the proxies rather than valuation. The main test is therefore the difference-in-differences $A^{U}-A^{L}$. The share recognised by week 1, $Recognition^{U\pm}(1)$, and a joint Wald test of $\beta^{U+}_{g}=\beta^{U-}_{g}$ for all six groups are also reported. Standard errors are Driscoll and Kraay (1998), and standard errors of ratios use the delta method.

Robustness: Model B (option and week fixed effects); excluding February to June 2020; defining the state by the sign of the listed equity market (the same for all options in a week); and comparing only the largest 10% of market falls with the largest 10% of rises. Every specification is estimated with equal weights and with June 2026 AUM weights.
""")

doc.add_heading('3. Results', level=1)
eu_p, eu_n = eP[('XU', 'p')], eP[('XU', 'n')]
wu_p, wu_n = wP[('XU', 'p')], wP[('XU', 'n')]
findings = [
    f'Negative and positive unlisted news are recognised with about the same overall delay. Equal-weighted, the '
    f'delayed share is {eu_p["share"][0]:.2f} for positive news and {eu_n["share"][0]:.2f} for negative news; the difference '
    f'is {eT["XU_dshare"][0]:+.3f} (standard error {eT["XU_dshare"][1]:.3f}, p = {p2(eT["XU_dshare"]):.2f}). AUM-weighted the '
    f'shares are {wu_p["share"][0]:.2f} and {wu_n["share"][0]:.2f} (difference {wT["XU_dshare"][0]:+.3f}, standard error '
    f'{wT["XU_dshare"][1]:.3f}, p = {p2(wT["XU_dshare"]):.2f}); a share slightly above one means that the same-week response '
    'is close to zero (Table 2).',
    f'The difference-in-differences against listed news is not significant in any of the ten specifications '
    f'(p-values {min(pvals_did):.2f} to {max(pvals_did):.2f}), and neither is the unlisted difference in delayed shares '
    f'(p-values {min(pvals_share):.2f} to {max(pvals_share):.2f}) (Table 3). H7 is not supported.',
    f'Within the delay, the point estimates suggest that negative news is recognised sooner: in the main equal-weighted '
    f'model {100 * eu_n["R1"][0]:.0f}% of the response to negative unlisted news is recognised by week 1, against '
    f'{100 * eu_p["R1"][0]:.0f}% for positive news, which is spread over weeks 2 to 8 (Figure 1). The difference is not '
    f'significant (p = {p2(eT["XU_dR1"]):.2f} equal-weighted, {p2(wT["XU_dR1"]):.2f} AUM-weighted); the joint Wald test of '
    f'identical lag profiles gives p = {1 - stats.chi2.cdf(eT["XU_wald"], 6):.2f} and '
    f'{1 - stats.chi2.cdf(wT["XU_wald"], 6):.2f} in Model A and p = {1 - stats.chi2.cdf(eB["XU_wald"], 6):.3f} and '
    f'{1 - stats.chi2.cdf(wB["XU_wald"], 6):.3f} with week fixed effects. This pattern, if confirmed, would be consistent '
    'with faster revaluation after market falls, but the present data cannot establish it.',
    f'Listed news, the placebo, responds slightly more strongly in the same week to negative news '
    f'({eP[("XL", "n")]["b0"][0]:.2f} against {eP[("XL", "p")]["b0"][0]:.2f} equal-weighted) and shows no delay in either '
    'direction, as expected for assets priced at market every day.',
    f'The test has low power. The standard error of the difference in delayed shares is about '
    f'{np.median(se_share):.2f}, so only a difference in delayed shares of about {MDE:.2f} or more '
    f'would be detected with 80% power at the 5% level. Splitting '
    'the news halves the information behind each coefficient, and large negative weeks are few and clustered (2020, 2022).']
for txt in findings:
    doc.add_paragraph(txt, style=style('List Bullet'))

add_table('Table 1', 'Positive and negative news in the estimation sample', t1, t1_note, [3.6, 3.2, 3.2, 3.2, 3.2], fs=8.5)
add_table('Table 2', 'Responses to positive and negative news: main specification', t2, t2_note,
          [4.8, 2.9, 2.9, 2.9, 2.9], fs=8.5, panel_rows=(0, len(ROWS) + 1))
add_table('Table 3', 'Tests of asymmetric recognition across specifications', t3, t3_note,
          [4.6, 2.5, 2.4, 2.5, 2.4, 2.0], fs=7.5)
add_figure('Figure 1', 'Recognition curves for positive and negative news (main specification)', FIG,
           'Cumulative share of the total response (through week 26) recognised by week h, from the point estimates of '
           'Table 2. Black: unlisted news; grey: listed news; solid: positive news; dashed: negative news. Shares above '
           'one reflect small negative coefficients at longer lags.')

doc.add_heading('4. Interpretation and use in the proposal', level=1)
for txt in [
    'The evidence does not show that unlisted assets are revalued more slowly after bad news than after good news. The '
    'overall delay is similar in both directions, and the difference relative to listed news is small and insignificant '
    'in every specification. The finding is a qualified null: the test cannot rule out moderate asymmetries.',
    'The weekly design measures the direction of each week\'s news, but unlisted assets are revalued mainly at quarter '
    'ends. A sharper test would compare the size of the quarter-end catch-up after quarters of falling and rising markets, '
    'which links naturally to the planned analysis of quarter-end and out-of-cycle revaluation dates.',
    'For the proposal, the question is best stated as a further analysis (or as a two-sided H7) with both competing '
    'explanations, rather than as a directional hypothesis.']:
    doc.add_paragraph(txt, style=style('List Bullet'))
doc.add_paragraph('References: APRA (2024b), Governance of unlisted asset valuation and liquidity risk management in '
                  'superannuation. Bollen, N. P. B. and Pool, V. K. (2008), Conditional return smoothing in the hedge fund '
                  'industry, Journal of Financial and Quantitative Analysis 43(2), 267-298. Driscoll, J. C. and Kraay, A. C. '
                  '(1998), Consistent covariance matrix estimation with spatially dependent panel data, Review of Economics '
                  'and Statistics 80(4), 549-560.')
fn = os.path.join(OUT, 'Asymmetric_recognition_results.docx')
doc.save(fn)

# ------------------------------------------------------------------ Excel
allp = []
for wgt in ('equal', 'aum'):
    for case, lab in CASES:
        a = R[(wgt,) + case]
        for key, m in a['parts'].items():
            allp.append(dict(weighting=wgt, specification=lab, news=key[0], state=key[1], n=a['n'],
                             **{f'{k}': m[k][0] for k in ('b0', 'D', 'T', 'share', 'R0', 'R1', 'R4', 'R8', 'R13', 'R26')},
                             **{f'{k}_se': m[k][1] for k in ('b0', 'D', 'T', 'share', 'R1', 'R4', 'R13')},
                             **{f'beta_{g}': b for g, b in zip(['0', '1', '2-4', '5-8', '9-13', '14-26'], m['b'])}))
curves = pd.DataFrame({'Weeks after the news': np.arange(27)})
for wgt in ('equal', 'aum'):
    for key in LAB:
        curves[f'{LAB[key]} ({wgt})'] = curve(R[(wgt,) + MAIN]['parts'][key]['b'])
with pd.ExcelWriter(os.path.join(OUT, 'asymmetric_recognition_tables.xlsx'), engine='openpyxl') as xw:
    for k, (title, t, note) in {'Table 1': ('Positive and negative news in the estimation sample', t1, t1_note),
                                'Table 2': ('Responses to positive and negative news: main specification', t2, t2_note),
                                'Table 3': ('Tests of asymmetric recognition across specifications', t3, t3_note)}.items():
        t.to_excel(xw, sheet_name=k, index=False, startrow=2)
        ws = xw.sheets[k]; ws['A1'] = f'{k}. {title}'; ws.cell(row=len(t) + 5, column=1, value='Note: ' + note)
    raw3.to_excel(xw, sheet_name='Tests (raw)', index=False)
    pd.DataFrame(allp).to_excel(xw, sheet_name='Estimates (raw)', index=False)
    curves.to_excel(xw, sheet_name='Recognition curves', index=False)
print('saved', fn)
print(t3.to_string())
