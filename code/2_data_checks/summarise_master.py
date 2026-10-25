"""
summarise_master.py -- summary statistics tables and figures for the option-quarter master dataset.

Run from the PhD_Topic3_Data folder:  python3 code/2_data_checks/summarise_master.py
Outputs: output/data_summary/  (Summary_statistics.docx, summary_tables.xlsx, fig*.png)

Main analysis sample = core_sample_dedup (multi-sector, multi-manager, accumulation options with complete
weights, one option per identical allocation per quarter). Weights are strategic benchmark weights.
"""
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import statsmodels.formula.api as smf

SRC = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
OUT = os.path.join(SRC, 'output', 'data_summary')
os.makedirs(OUT, exist_ok=True)

m = pd.read_csv(os.path.join(SRC, 'data', 'derived', 'option_quarter_master.csv'), low_memory=False,
                parse_dates=['saa_date', 'fy_end', 'last_major_merger_fy'])
fy = pd.read_csv(os.path.join(SRC, 'data', 'derived', 'fund_year.csv'), parse_dates=['fy_end'])
TYPES = ['Industry', 'Public Sector', 'Retail', 'Corporate']
FIRST, LAST = m['saa_date'].min(), m['saa_date'].max()
m['growth'] = m[['w_EQ_AU_L', 'w_EQ_INT_L', 'w_EQ_OTHER', 'w_PE_U', 'w_PROP_AU_L', 'w_PROP_INT_L', 'w_PROP_U_AU',
                 'w_PROP_U_INT', 'w_PROP_U_NSDOM', 'w_PROP_NS', 'w_PROP_OTHER', 'w_INFRA_L', 'w_INFRA_U', 'w_INFRA_NS',
                 'w_INFRA_OTHER', 'w_ALT_U', 'w_ALT_OTHER']].sum(axis=1)
m['w_prop_l'] = m['w_PROP_AU_L'] + m['w_PROP_INT_L']
m['w_defensive'] = m['w_FI'] + m['w_CASH']
S = m[m['core_sample_dedup']].copy()          # main analysis sample
tables = {}

# ----------------------------------------------------------------------------- Table 1: sample construction
steps = [
    ('All options in APRA Historical SAA', pd.Series(True, index=m.index)),
    ('Multi-sector options', m['option_category'].eq('Multi Sector')),
]
steps.append(('Multi-manager', steps[-1][1] & m['option_type'].eq('Multi Manager')))
steps.append(('Excluding pension options', steps[-1][1] & ~m['is_pension']))
steps.append(('Excluding incomplete weights and bad-weight funds', steps[-1][1] & ~m['weights_incomplete'] & ~m['bad_weight_fund']))
steps.append(('Excluding specialist options (core sample)', m['core_sample']))
steps.append(('One option per identical allocation (analysis sample)', m['core_sample_dedup']))
rows = []
for name, mask in steps:
    d = m[mask]; l = d[d['saa_date'] == LAST]; f = d[d['saa_date'] == FIRST]
    rows.append([name, len(d), d['option_id'].nunique(), d['abn'].nunique(), len(f), len(l), l['abn'].nunique()])
t1 = pd.DataFrame(rows, columns=['Step', 'Option-quarters', 'Options', 'Funds',
                                 f'Options, {FIRST:%b %Y}', f'Options, {LAST:%b %Y}', f'Funds, {LAST:%b %Y}'])
tables['Table 1'] = ('Sample construction', t1,
                     'Option-quarters from the APRA Quarterly Superannuation Product Statistics Historical SAA file, '
                     f'{FIRST:%B %Y} to {LAST:%B %Y}. Pension options are offered only in the retirement phase or named as pension '
                     'or retirement income options. Incomplete weights: weights excluding the currency overlay more than '
                     '2 percentage points from 100%. Bad-weight funds: Avanteos, Netwealth and Wealth Personal. Specialist options: '
                     'capital-guaranteed or life products, options with property, infrastructure, credit or alternatives of at least '
                     '40%, cash or fixed income of at least 80%, or unlisted assets only (e.g. HESTA Property and Infrastructure, '
                     'ART Unlisted Assets).')

# ----------------------------------------------------------------------------- Table 2: composition by fund type
rows = []
for q in (FIRST, LAST):
    d = S[S['saa_date'] == q]
    for t in TYPES + ['All']:
        x = d if t == 'All' else d[d['fund_type_fixed'] == t]
        if len(x) == 0:
            continue
        rows.append([f'{q:%b %Y}', t, x['abn'].nunique(), len(x), 100 * x['w_unlisted_core'].mean(),
                     100 * x['w_unlisted_core'].median(), 100 * (x['w_unlisted_core'] > 0).mean(),
                     100 * x['w_unlisted_core'].max()])
t2 = pd.DataFrame(rows, columns=['Quarter', 'Fund type', 'Funds', 'Options', 'Mean unlisted (%)',
                                 'Median unlisted (%)', 'Options with any unlisted (%)', 'Maximum unlisted (%)'])
tables['Table 2'] = ('Unlisted exposure by fund type, first and last quarter', t2,
                     'Analysis sample. Unlisted = w_unlisted_core: unlisted equity, property, infrastructure and credit '
                     'as a share of the option. Fund type is the June 2026 APRA classification.')

# ----------------------------------------------------------------------------- Table 3: descriptive statistics
V = [('w_unlisted_core', 'Unlisted assets (core)', 100), ('w_PE_U', '  Unlisted equity', 100),
     ('w_prop_u', '  Unlisted property', 100), ('w_INFRA_U', '  Unlisted infrastructure', 100),
     ('w_CREDIT_U', '  Unlisted credit', 100), ('w_ALT_U', 'Unlisted alternatives', 100),
     ('w_notspecified_real', 'Property/infrastructure, listing not specified', 100),
     ('w_listed_growth', 'Listed growth assets', 100), ('w_EQ_AU_L', '  Listed Australian equity', 100),
     ('w_EQ_INT_L', '  Listed international equity', 100), ('w_prop_l', '  Listed property', 100),
     ('w_INFRA_L', '  Listed infrastructure', 100), ('w_FI', 'Fixed income', 100), ('w_CASH', 'Cash', 100),
     ('growth', 'Growth assets (all)', 100), ('w_intl', 'International assets', 100),
     ('hedge_ratio_intl', 'Hedge ratio of international assets', 100),
     ('apra_return_measure', 'APRA quarterly return measure', 100), ('apra_vol_5y', 'APRA 5-year volatility', 100)]
rows = []
for v, lab, k in V:
    x = S[v].dropna() * k
    rows.append([lab, len(x), x.mean(), x.std(), x.quantile(.10), x.quantile(.25), x.median(), x.quantile(.75), x.quantile(.90)])
t3 = pd.DataFrame(rows, columns=['Variable (%)', 'N', 'Mean', 'SD', 'P10', 'P25', 'Median', 'P75', 'P90'])
tables['Table 3'] = ('Descriptive statistics of allocation variables', t3,
                     f'Analysis sample, all option-quarters {FIRST:%b %Y}-{LAST:%b %Y}. All variables in percent. Weights are '
                     'strategic benchmark weights. Growth assets: equity, property, infrastructure and alternatives, listed or '
                     'unlisted. The APRA return and volatility measures come from the QSPS Historical performance file; their '
                     'exact definitions should be confirmed with APRA (the return measure behaves like a quarterly net return).')

# ----------------------------------------------------------------------------- Table 4: variation in unlisted exposure
S = S.sort_values(['option_id', 'saa_date'])
S['lag_unl'] = S.groupby('option_id')['w_unlisted_core'].shift()
opt_mean = S.groupby('option_id')['w_unlisted_core'].transform('mean')
within = (S['w_unlisted_core'] - opt_mean)
chg = S['w_unlisted_core'] - S['lag_unl']
per = S.dropna(subset=['lag_unl'])
rows = [
    ['Overall SD (pp)', 100 * S['w_unlisted_core'].std()],
    ['Between-option SD (pp)', 100 * S.groupby('option_id')['w_unlisted_core'].mean().std()],
    ['Within-option SD (pp)', 100 * within.std()],
    ['Share of variance between options (%)', 100 * (1 - within.var() / S['w_unlisted_core'].var())],
    ['Cross-sectional SD within a quarter, average (pp)', 100 * S.groupby('saa_date')['w_unlisted_core'].std().mean()],
    ['Quarter-to-quarter correlation', per['w_unlisted_core'].corr(per['lag_unl'])],
    ['Option-quarters with any change (%)', 100 * (chg.abs() > 1e-9).mean()],
    ['Changes larger than 10 pp (count)', int((chg.abs() > 0.10).sum())],
    ['  of which likely reclassification (count)', int(S['reclass_jump'].sum())],
    ['Options', S['option_id'].nunique()],
    ['Option-quarters', len(S)],
]
t4 = pd.DataFrame([[a, f'{int(b):,}' if isinstance(b, (int, np.integer)) else f'{b:.2f}'] for a, b in rows],
                  columns=['Statistic', 'Value'])
tables['Table 4'] = ('Variation in unlisted exposure', t4,
                     'Analysis sample. pp = percentage points. Between-option SD: standard deviation of option means. '
                     'Within-option SD: standard deviation of deviations from option means. The identification of the '
                     'delayed-response model relies on cross-sectional differences in exposure within the same week, '
                     'summarised by the cross-sectional SD within a quarter.')

# ----------------------------------------------------------------------------- Table 5: fund characteristics
core_funds = m.loc[m['core_sample'], ['abn', 'fund_type_fixed', 'last_major_merger_fy']].drop_duplicates('abn')
f = fy.merge(core_funds, on='abn', how='inner')
f = f[(f['fy_end'] >= '2014-06-30') & ~(f['last_major_merger_fy'].notna() & (f['fy_end'] <= f['last_major_merger_fy']))]
rows = []
for t in TYPES + ['All']:
    x = f if t == 'All' else f[f['fund_type_fixed'] == t]
    if len(x) == 0:
        continue
    rows.append([t, x['abn'].nunique(), len(x), x['total_assets_k'].mean() / 1e6, x['total_assets_k'].median() / 1e6,
                 x['member_accounts'].median() / 1e3, 100 * x['net_flow_ratio_ex_sft_w'].median(),
                 100 * x['ret_1y_fund'].mean(), 100 * x['ret_1y_fund'].std()])
t5 = pd.DataFrame(rows, columns=['Fund type', 'Funds', 'Fund-years', 'Total assets, mean ($bn)', 'Total assets, median ($bn)',
                                 "Member accounts, median ('000)", 'Net flows excl. mergers, median (% of assets)',
                                 'One-year return, mean (%)', 'One-year return, SD (%)'])
tables['Table 5'] = ('Fund characteristics', t5,
                     'Funds with at least one option in the core sample; fund-years FY2014-FY2025 from APRA annual fund-level '
                     'statistics, excluding years up to a major merger (successor fund transfers in of at least 25% of '
                     'opening assets). Net flows exclude successor fund transfers and are winsorised at the 1st and 99th '
                     'percentiles. Returns are whole-of-fund one-year returns reported by APRA.')

# ----------------------------------------------------------------------------- Table 6: reported volatility and unlisted exposure
V6 = S[(S['saa_date'] == LAST)].dropna(subset=['apra_vol_5y']).copy()
V6['vol'] = 100 * V6['apra_vol_5y']; V6['unl'] = 100 * V6['w_unlisted_core']; V6['gr'] = 100 * V6['growth']
res = []
for name, fml in [('(1)', 'vol ~ gr'), ('(2)', 'vol ~ gr + unl'), ('(3)', 'vol ~ gr + unl + C(fund_type_fixed)')]:
    r = smf.ols(fml, data=V6).fit(cov_type='cluster', cov_kwds={'groups': pd.factorize(V6['abn'])[0]})
    res.append((name, r))
rows = []
for lab, v in [('Growth assets (%)', 'gr'), ('Unlisted assets (%)', 'unl')]:
    rows.append([lab] + [f"{r.params[v]:.3f}{'***' if r.pvalues[v] < .01 else '**' if r.pvalues[v] < .05 else '*' if r.pvalues[v] < .1 else ''}"
                         if v in r.params else '' for _, r in res])
    rows.append([''] + [f"({r.bse[v]:.3f})" if v in r.params else '' for _, r in res])
rows.append(['Fund-type effects'] + ['No', 'No', 'Yes'])
rows.append(['Observations'] + [f'{int(r.nobs)}' for _, r in res])
rows.append(['R-squared'] + [f'{r.rsquared:.3f}' for _, r in res])
t6 = pd.DataFrame(rows, columns=['', '(1)', '(2)', '(3)'])
b_unl = res[1][1].params['unl']
tables['Table 6'] = ('Reported volatility and unlisted exposure (preliminary)', t6,
                     f'OLS of the APRA 5-year volatility measure (%) on the growth-asset share and the unlisted share (%), '
                     f'analysis sample, {LAST:%B %Y}. Standard errors clustered by fund in parentheses. *** p<0.01, ** p<0.05, '
                     '* p<0.1. Descriptive only: a negative unlisted coefficient means options with more unlisted assets report '
                     'lower volatility for the same growth share, which is consistent with, but does not prove, smoothed valuation.')

# ----------------------------------------------------------------------------- Table 7: real change vs composition
def decompose(d, v, q0, q1):
    a = d[d['saa_date'] == q0].set_index('option_id')[v]; b = d[d['saa_date'] == q1].set_index('option_id')[v]
    both = a.index.intersection(b.index)
    total = b.mean() - a.mean()
    within = (b[both] - a[both]).sum() / len(b)
    return len(a), len(b), len(both), len(b.index.difference(a.index)), len(a.index.difference(b.index)), \
        100 * a.mean(), 100 * b.mean(), 100 * total, 100 * within, 100 * (total - within)


rows = []
junes = [q for q in sorted(S['saa_date'].unique()) if pd.Timestamp(q).month == 6]
for q0, q1 in zip(junes[:-1], junes[1:]):
    for t in ['All', 'Industry', 'Public Sector', 'Retail']:
        d = S if t == 'All' else S[S['fund_type_fixed'] == t]
        n0, n1, nb, ne, nx, m0, m1, tot, wi, comp = decompose(d, 'w_unlisted_core', q0, q1)
        rows.append([f'{pd.Timestamp(q0):%Y}-{pd.Timestamp(q1):%y}', t, n0, n1, ne, nx, m0, m1, tot, wi, comp])
t7 = pd.DataFrame(rows, columns=['Year to June', 'Fund type', 'Options, start', 'Options, end', 'Entrants', 'Exits',
                                 'Mean, start (%)', 'Mean, end (%)', 'Change (pp)', 'Within options (pp)', 'Composition (pp)'])
t7_all = t7[t7['Fund type'] == 'All'].drop(columns='Fund type')
tables['Table 7'] = ('Change in mean unlisted share: real reallocation versus sample composition', t7_all,
                     'Analysis sample, all fund types. Change = change in the mean unlisted share from June to June. Within options '
                     '= contribution of allocation changes in options present in both years (sum of their changes divided by the '
                     'number of options at the end). Composition = the remainder, due to options entering or leaving the sample. '
                     'The same decomposition by fund type is in summary_tables.xlsx (sheet Table 7 by type).')

# balanced panel: options present in every quarter
nq = S['saa_date'].nunique()
bal_ids = S.groupby('option_id')['saa_date'].nunique()
bal_ids = bal_ids[bal_ids == nq].index
B = S[S['option_id'].isin(bal_ids)]

# ----------------------------------------------------------------------------- figures (black and white)
plt.rcParams.update({'font.family': 'serif', 'font.serif': ['Times New Roman', 'Liberation Serif', 'DejaVu Serif'],
                     'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False,
                     'axes.edgecolor': '0.3', 'axes.linewidth': 0.8, 'xtick.color': '0.2', 'ytick.color': '0.2',
                     'axes.grid': True, 'axes.axisbelow': True, 'grid.color': '0.9', 'grid.linewidth': 0.6, 'legend.frameon': False,
                     'savefig.dpi': 300, 'savefig.bbox': 'tight'})
STY = {'Industry': dict(color='0.0', ls='-', marker='o'), 'Public Sector': dict(color='0.35', ls='--', marker='s'),
       'Retail': dict(color='0.55', ls=':', marker='^'), 'Corporate': dict(color='0.7', ls='-.', marker='D')}
figs = {}


def pct(ax):
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f'{v:.0f}%'))


# Figure 1: sample size over time (two panels, no dual axis)
g = m.groupby('saa_date').agg(core=('core_sample', 'sum'), dedup=('core_sample_dedup', 'sum'))
gf = m[m['core_sample']].groupby('saa_date')['abn'].nunique()
fig, ax = plt.subplots(1, 2, figsize=(7.2, 2.8))
ax[0].plot(g.index, g['core'], color='0.0', lw=1.5, label='Core sample')
ax[0].plot(g.index, g['dedup'], color='0.45', lw=1.5, ls='--', label='One per identical allocation')
ax[0].set_title('Options', loc='left', fontsize=10); ax[0].legend(loc='upper left', fontsize=8.5)
ax[1].plot(gf.index, gf.values, color='0.0', lw=1.5)
ax[1].set_title('Funds', loc='left', fontsize=10)
for a in ax:
    a.set_ylim(bottom=0)
fig.tight_layout(); p = os.path.join(OUT, 'fig1_sample_over_time.png'); fig.savefig(p); plt.close(fig)
figs['Figure 1'] = ('Sample size over time', p, 'Number of options and funds in the core sample each quarter. '
                    'The dashed line keeps one option per identical allocation (analysis sample).')

# Figure 2: mean unlisted share by fund type
fig, ax = plt.subplots(figsize=(7.2, 3.6))
for t in ['Industry', 'Public Sector', 'Retail']:
    s = S[S['fund_type_fixed'] == t].groupby('saa_date')['w_unlisted_core'].mean() * 100
    st = STY[t]
    ax.plot(s.index, s.values, color=st['color'], ls=st['ls'], lw=1.6, marker=st['marker'], markersize=3.5,
            markevery=4, label=t)
    ax.annotate(f'{t} {s.iloc[-1]:.1f}%', (s.index[-1], s.iloc[-1]), xytext=(6, 0), textcoords='offset points',
                va='center', fontsize=8.5, annotation_clip=False)
pct(ax); ax.set_ylim(bottom=0); ax.set_ylabel('Mean unlisted share of option')
ax.legend(loc='upper left', fontsize=8.5)
ax.set_xlim(S['saa_date'].min() - pd.Timedelta(days=90), S['saa_date'].max() + pd.Timedelta(days=60))
p = os.path.join(OUT, 'fig2_unlisted_by_fund_type.png'); fig.savefig(p); plt.close(fig)
figs['Figure 2'] = ('Unlisted exposure by fund type', p, 'Mean of w_unlisted_core (unlisted equity, property, '
                    'infrastructure and credit) across options in the analysis sample, by June 2026 fund type. '
                    'Corporate funds (four options) are omitted.')

# Figure 3: composition of the unlisted share over time (parts of a whole -> stacked)
comp = S.groupby('saa_date')[['w_prop_u', 'w_INFRA_U', 'w_PE_U', 'w_CREDIT_U']].mean() * 100
fig, ax = plt.subplots(figsize=(7.2, 3.4))
fills = [('w_prop_u', 'Property', '0.25', ''), ('w_INFRA_U', 'Infrastructure', '0.5', ''),
         ('w_PE_U', 'Private equity', '0.75', ''), ('w_CREDIT_U', 'Private credit', '1.0', '////')]
base = np.zeros(len(comp))
for col, lab, shade, hatch in fills:
    ax.fill_between(comp.index, base, base + comp[col].values, facecolor=shade, edgecolor='0.0', lw=0.4,
                    hatch=hatch, label=lab)
    base = base + comp[col].values
pct(ax); ax.set_ylim(bottom=0); ax.set_ylabel('Mean share of option')
h, lbl = ax.get_legend_handles_labels(); ax.legend(h[::-1], lbl[::-1], loc='upper left', fontsize=8.5)
p = os.path.join(OUT, 'fig3_unlisted_composition.png'); fig.savefig(p); plt.close(fig)
figs['Figure 3'] = ('Composition of unlisted exposure', p, 'Mean weight of each unlisted asset class across options '
                    'in the analysis sample; the four layers add up to the mean of w_unlisted_core.')

# Figure 2b / 3b: the same on the balanced panel (fixed set of options)
fig, ax = plt.subplots(1, 2, figsize=(7.2, 3.8), sharey=False)
for t in ['Industry', 'Public Sector', 'Retail']:
    s_ = B[B['fund_type_fixed'] == t].groupby('saa_date')['w_unlisted_core'].mean() * 100
    if len(s_) == 0:
        continue
    st = STY[t]; n_t = B.loc[B['fund_type_fixed'] == t, 'option_id'].nunique()
    ax[0].plot(s_.index, s_.values, color=st['color'], ls=st['ls'], lw=1.5, marker=st['marker'], markersize=3,
               markevery=4, label=f'{t} (n={n_t})')
ax[0].set_title('Mean unlisted share by fund type', loc='left', fontsize=10); pct(ax[0]); ax[0].set_ylim(bottom=0)
ax[0].legend(loc='upper center', bbox_to_anchor=(0.5, -0.12), ncol=2, fontsize=7.5)
compb = B.groupby('saa_date')[['w_prop_u', 'w_INFRA_U', 'w_PE_U', 'w_CREDIT_U']].mean() * 100
base = np.zeros(len(compb))
for col_, lab, shade, hatch in fills:
    ax[1].fill_between(compb.index, base, base + compb[col_].values, facecolor=shade, edgecolor='0.0', lw=0.4,
                       hatch=hatch, label=lab)
    base = base + compb[col_].values
ax[1].set_title('Composition, all fund types', loc='left', fontsize=10); pct(ax[1]); ax[1].set_ylim(bottom=0)
h, lbl = ax[1].get_legend_handles_labels(); ax[1].legend(h[::-1], lbl[::-1], loc='upper center', bbox_to_anchor=(0.5, -0.12), ncol=2, fontsize=7.5)
for a in ax:
    a.tick_params(axis='x', labelsize=8)
fig.tight_layout(); p = os.path.join(OUT, 'fig2b_balanced_panel.png'); fig.savefig(p); plt.close(fig)
figs['Figure 2b'] = ('Unlisted exposure in a balanced panel', p, f'Options in the analysis sample observed in every quarter '
                     f'{FIRST:%B %Y} to {LAST:%B %Y} ({len(bal_ids)} options). With a fixed set of options, changes reflect '
                     'reallocation only; compare with Figures 2 and 3, whose averages also move when options enter the sample.')

# Figure 4: cross-sectional dispersion (identifying variation) over time
q = S.groupby('saa_date')['w_unlisted_core'].quantile([.1, .25, .5, .75, .9]).unstack() * 100
fig, ax = plt.subplots(figsize=(7.2, 3.4))
ax.fill_between(q.index, q[0.1], q[0.9], facecolor='0.88', edgecolor='none', label='10th to 90th percentile')
ax.fill_between(q.index, q[0.25], q[0.75], facecolor='0.65', edgecolor='none', label='25th to 75th percentile')
ax.plot(q.index, q[0.5], color='0.0', lw=1.6, label='Median')
pct(ax); ax.set_ylim(bottom=0); ax.set_ylabel('Unlisted share of option')
ax.legend(loc='upper left', fontsize=8.5)
p = os.path.join(OUT, 'fig4_unlisted_dispersion.png'); fig.savefig(p); plt.close(fig)
figs['Figure 4'] = ('Cross-sectional spread of unlisted exposure', p, 'Percentiles of w_unlisted_core across options '
                    'in the analysis sample each quarter. This spread across options in the same week is the variation '
                    'that identifies the delayed-response model.')

# Figure 5: distribution by fund type, last quarter (one measure across groups -> box plot with points)
d = S[S['saa_date'] == LAST]
fig, ax = plt.subplots(figsize=(7.2, 3.2))
grp = ['Industry', 'Public Sector', 'Retail']
data = [d.loc[d['fund_type_fixed'] == t, 'w_unlisted_core'].values * 100 for t in grp]
ax.boxplot(data, vert=False, widths=0.5, showfliers=False, medianprops=dict(color='0.0', lw=1.6),
           boxprops=dict(color='0.2'), whiskerprops=dict(color='0.2'), capprops=dict(color='0.2'))
rng = np.random.default_rng(1)
for i, x in enumerate(data):
    ax.scatter(x, i + 1 + rng.uniform(-0.18, 0.18, len(x)), s=7, facecolor='none', edgecolor='0.35', lw=0.6)
ax.set_yticks(range(1, len(grp) + 1), [f'{t} (n={len(x)})' for t, x in zip(grp, data)])
ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f'{v:.0f}%'))
ax.set_xlabel(f'Unlisted share of option, {LAST:%B %Y}'); ax.grid(axis='y', visible=False)
p = os.path.join(OUT, 'fig5_unlisted_distribution.png'); fig.savefig(p); plt.close(fig)
figs['Figure 5'] = ('Distribution of unlisted exposure across options', p, f'Analysis sample, {LAST:%B %Y}. Boxes show the '
                    'interquartile range and median; whiskers extend to 1.5 times the interquartile range; each circle is an option.')

# Figure 6: reported volatility vs growth share, low vs high unlisted (binned means)
V6['unl_group'] = np.where(V6['unl'] >= V6['unl'].quantile(2 / 3), 'High unlisted (top third)',
                           np.where(V6['unl'] <= V6['unl'].quantile(1 / 3), 'Low unlisted (bottom third)', 'Middle'))
bins = np.arange(0, 110, 10)
V6['gbin'] = pd.cut(V6['gr'], bins, labels=bins[:-1] + 5)
fig, ax = plt.subplots(figsize=(7.2, 3.6))
for lab, st in [('Low unlisted (bottom third)', dict(color='0.55', ls='--', marker='s')),
                ('High unlisted (top third)', dict(color='0.0', ls='-', marker='o'))]:
    b = V6[V6['unl_group'] == lab].groupby('gbin', observed=True)['vol'].agg(['mean', 'size'])
    b = b[b['size'] >= 5]
    ax.plot(b.index.astype(float), b['mean'], lw=1.6, markersize=5, label=lab, **st)
ax.set_xlabel('Growth-asset share of option (bins of 10 pp)')
ax.set_ylabel('Reported 5-year volatility')
ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f'{v:.0f}%'))
pct(ax); ax.legend(loc='upper left', fontsize=8.5)
p = os.path.join(OUT, 'fig6_volatility_vs_growth.png'); fig.savefig(p); plt.close(fig)
figs['Figure 6'] = ('Reported volatility, growth share and unlisted exposure', p, f'Mean APRA 5-year volatility measure by '
                    f'growth-asset bin for options in the bottom and top third of unlisted exposure, analysis sample, '
                    f'{LAST:%B %Y}; bins with fewer than five options are omitted. Holding growth share fixed, a 10 pp higher '
                    f'unlisted share is associated with {10 * b_unl:.2f} pp {"lower" if b_unl < 0 else "higher"} reported volatility '
                    '(Table 6, column 2). Preliminary and descriptive.')

# ----------------------------------------------------------------------------- Excel tables
with pd.ExcelWriter(os.path.join(OUT, 'summary_tables.xlsx'), engine='openpyxl') as xw:
    for k, (title, t, note) in tables.items():
        t.to_excel(xw, sheet_name=k, index=False, startrow=2)
        ws = xw.sheets[k]; ws['A1'] = f'{k}. {title}'; ws.cell(row=len(t) + 5, column=1, value='Note: ' + note)
    t7.to_excel(xw, sheet_name='Table 7 by type', index=False)

# ----------------------------------------------------------------------------- Word document (Times New Roman, black and white)
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for a in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(s, a, Cm(2.3))
for st in doc.styles:
    try:
        st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a in list(rf.attrib):
            del rf.attrib[a]
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), 'Times New Roman')
    except Exception:
        pass
doc.styles['Normal'].font.size = Pt(11)
for h, sz in (('Title', 16), ('Heading 1', 13), ('Heading 2', 11.5)):
    doc.styles[h].font.size = Pt(sz); doc.styles[h].font.bold = True


def border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:color'), '000000')
    b.append(e)


def fmt(v, colname):
    if isinstance(v, str):
        return v
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ''
    if isinstance(v, (int, np.integer)):
        return f'{int(v):,}'
    return f'{v:,.2f}'


def add_table(num, title, t, note):
    p = doc.add_paragraph(); r = p.add_run(f'{num}. {title}'); r.bold = True
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
    tb = doc.add_table(rows=len(t) + 1, cols=t.shape[1]); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb.autofit = False
    first = {'Table 1': 5.6, 'Table 3': 4.6, 'Table 4': 6.0, 'Table 6': 4.0, 'Table 7': 1.8}.get(num, 2.6)
    other = (16.4 - first) / max(t.shape[1] - 1, 1)
    widths = [first] + [other] * (t.shape[1] - 1)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 567)))
    tblPr = tb._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    for j, w in enumerate(widths):
        for c in tb.columns[j].cells:
            c.width = Cm(w)
    for j, c in enumerate(t.columns):
        tb.cell(0, j).text = str(c)
    for i, row in enumerate(t.itertuples(index=False)):
        for j, v in enumerate(row):
            tb.cell(i + 1, j).text = fmt(v, t.columns[j])
    n = len(t) + 1
    for i in range(n):
        for j in range(t.shape[1]):
            c = tb.cell(i, j)
            if i == 0:
                border(c, 'top', 12); border(c, 'bottom', 6)
            if i == n - 1:
                border(c, 'bottom', 12)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.space_before = Pt(1)
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                pp.paragraph_format.keep_with_next = i < n - 1      # keep each table on one page
                for rr in pp.runs:
                    rr.font.size = Pt(9); rr.font.bold = (i == 0)
    q = doc.add_paragraph(); rn = q.add_run('Note: ' + note); rn.font.size = Pt(9); rn.italic = False


def add_figure(num, title, path, note):
    p = doc.add_paragraph(); r = p.add_run(f'{num}. {title}'); r.bold = True
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
    doc.add_picture(path, width=Cm(16)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    q = doc.add_paragraph(); rn = q.add_run('Note: ' + note); rn.font.size = Pt(9)


doc.add_heading('Summary statistics: option-quarter master dataset', level=0)
doc.add_paragraph(f'Minh Huong · data {FIRST:%B %Y} to {LAST:%B %Y} · generated by code/2_data_checks/summarise_master.py')
d_last = S[S['saa_date'] == LAST]
intro = (f'The analysis sample contains {S["option_id"].nunique():,} multi-sector, multi-manager accumulation options from '
         f'{S["abn"].nunique()} APRA-regulated funds, observed quarterly from {FIRST:%B %Y} to {LAST:%B %Y} '
         f'({len(S):,} option-quarters). In {LAST:%B %Y} the average option held '
         f'{100 * d_last["w_unlisted_core"].mean():.1f}% in unlisted assets (median {100 * d_last["w_unlisted_core"].median():.1f}%), '
         f'with a 10th-90th percentile range of {100 * d_last["w_unlisted_core"].quantile(.1):.1f}% to '
         f'{100 * d_last["w_unlisted_core"].quantile(.9):.1f}%. Industry and public sector funds hold far more unlisted '
         'assets than retail funds. All weights are strategic benchmark weights from APRA.')
doc.add_paragraph(intro)
doc.add_heading('Tables', level=1)
for k, (title, t, note) in tables.items():
    add_table(k, title, t, note)
doc.add_page_break()
doc.add_heading('Figures', level=1)
for k, (title, path, note) in figs.items():
    add_figure(k, title, path, note)
doc.save(os.path.join(OUT, 'Summary_statistics.docx'))
print('done:', len(tables), 'tables,', len(figs), 'figures ->', OUT)
print(t2.round(1).to_string()); print(t4.round(3).to_string()); print(t6.to_string())
